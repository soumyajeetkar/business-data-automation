import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.utils import get_column_letter


def generate_report(
    df,
    metrics,
    product_sales,
    category_sales,
    monthly_sales,
    output_file
):

    summary = pd.DataFrame({
        "Metric": [
            "Total Revenue",
            "Total Quantity Sold",
            "Total Orders",
            "Average Order Value",
            "Best Product",
            "Best Product Revenue",
            "Best Category",
            "Best Category Revenue"
        ],

        "Value": [
            metrics["total_sales"],
            metrics["total_quantity"],
            metrics["total_orders"],
            metrics["average_order_value"],
            product_sales.index[0],
            product_sales.iloc[0],
            category_sales.index[0],
            category_sales.iloc[0]
        ]
    })


    with pd.ExcelWriter(
        output_file,
        engine="openpyxl"
    ) as writer:

        summary.to_excel(
            writer,
            sheet_name="Summary",
            index=False
        )

        df.to_excel(
            writer,
            sheet_name="Sales Data",
            index=False
        )

        product_sales.reset_index().to_excel(
            writer,
            sheet_name="Product Analysis",
            index=False
        )

        category_sales.reset_index().to_excel(
            writer,
            sheet_name="Category Analysis",
            index=False
        )

        monthly_sales.to_excel(
            writer,
            sheet_name="Monthly Analysis",
            index=False
        )


    workbook = load_workbook(output_file)


    header_fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78"
    )

    header_font = Font(
        color="FFFFFF",
        bold=True
    )


    for worksheet in workbook.worksheets:

        for cell in worksheet[1]:

            cell.fill = header_fill
            cell.font = header_font

            cell.alignment = Alignment(
                horizontal="center"
            )

        worksheet.freeze_panes = "A2"

        worksheet.auto_filter.ref = (
            worksheet.dimensions
        )


        for column_cells in worksheet.columns:

            max_length = 0

            column_letter = get_column_letter(
                column_cells[0].column
            )

            for cell in column_cells:

                if cell.value is not None:

                    max_length = max(
                        max_length,
                        len(str(cell.value))
                    )

            worksheet.column_dimensions[
                column_letter
            ].width = min(
                max_length + 3,
                35
            )


    # Product chart

    product_sheet = workbook["Product Analysis"]

    chart = BarChart()

    chart.title = "Revenue by Product"

    data = Reference(
        product_sheet,
        min_col=2,
        min_row=1,
        max_row=product_sheet.max_row
    )

    categories = Reference(
        product_sheet,
        min_col=1,
        min_row=2,
        max_row=product_sheet.max_row
    )

    chart.add_data(
        data,
        titles_from_data=True
    )

    chart.set_categories(categories)

    product_sheet.add_chart(
        chart,
        "D2"
    )


    # Category chart

    category_sheet = workbook[
        "Category Analysis"
    ]

    category_chart = BarChart()

    category_chart.title = "Revenue by Category"

    data = Reference(
        category_sheet,
        min_col=2,
        min_row=1,
        max_row=category_sheet.max_row
    )

    categories = Reference(
        category_sheet,
        min_col=1,
        min_row=2,
        max_row=category_sheet.max_row
    )

    category_chart.add_data(
        data,
        titles_from_data=True
    )

    category_chart.set_categories(
        categories
    )

    category_sheet.add_chart(
        category_chart,
        "D2"
    )


    # Monthly chart

    monthly_sheet = workbook[
        "Monthly Analysis"
    ]

    monthly_chart = LineChart()

    monthly_chart.title = "Monthly Revenue"

    data = Reference(
        monthly_sheet,
        min_col=2,
        min_row=1,
        max_row=monthly_sheet.max_row
    )

    categories = Reference(
        monthly_sheet,
        min_col=1,
        min_row=2,
        max_row=monthly_sheet.max_row
    )

    monthly_chart.add_data(
        data,
        titles_from_data=True
    )

    monthly_chart.set_categories(
        categories
    )

    monthly_sheet.add_chart(
        monthly_chart,
        "D2"
    )


    workbook.save(output_file)